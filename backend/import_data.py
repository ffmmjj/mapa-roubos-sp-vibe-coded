#!/usr/bin/env python3
"""Import XLSX data into SQLite database."""
import openpyxl
import sqlite3
import sys
import os
from datetime import datetime

DB_PATH = os.path.join(os.path.dirname(__file__), "crimes.db")
XLSX_PATH = os.path.join(os.path.dirname(__file__), "..", "CelularesSubtraidos_2025.xlsx")

# Columns we want to import (in order)
WANTED_COLS = [
    'ID_DELEGACIA', 'NOME_DEPARTAMENTO', 'NOME_SECCIONAL', 'NOME_DELEGACIA',
    'NOME_MUNICIPIO', 'ANO_BO', 'NUM_BO', 'VERSAO',
    'NOME_DEPARTAMENTO_CIRC', 'NOME_SECCIONAL_CIRC', 'NOME_DELEGACIA_CIRC',
    'NOME_MUNICIPIO_CIRC',
    'DATA_OCORRENCIA_BO', 'HORA_OCORRENCIA', 'DESCRICAO_APRESENTACAO',
    'DATAHORA_REGISTRO_BO', 'DATA_COMUNICACAO_BO',
    'DESCR_PERIODO', 'AUTORIA_BO', 'FLAG_FLAGRANTE', 'FLAG_STATUS',
    'RUBRICA', 'DESCR_CONDUTA', 'DESCR_TIPOLOCAL', 'DESCR_SUBTIPOLOCAL',
    'CIDADE', 'BAIRRO', 'CEP', 'LOGRADOURO', 'NUMERO_LOGRADOURO',
    'LATITUDE', 'LONGITUDE', 'MARCA_OBJETO',
    'FLAG_BLOQUEIO', 'FLAG_DESBLOQUEIO', 'MES_REGISTRO_BO', 'ANO_REGISTRO_BO'
]

def create_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL")
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS crimes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            id_delegacia INTEGER,
            nome_departamento TEXT,
            nome_seccional TEXT,
            nome_delegacia TEXT,
            nome_municipio TEXT,
            ano_bo INTEGER,
            num_bo TEXT,
            versao INTEGER,
            nome_departamento_circ TEXT,
            nome_seccional_circ TEXT,
            nome_delegacia_circ TEXT,
            nome_municipio_circ TEXT,
            data_ocorrencia_bo TEXT,
            hora_ocorrencia TEXT,
            descricao_apresentacao TEXT,
            datahora_registro_bo TEXT,
            data_comunicacao_bo TEXT,
            descr_periodo TEXT,
            autoria_bo TEXT,
            flag_flagrante TEXT,
            flag_status TEXT,
            rubrica TEXT,
            descr_conduta TEXT,
            descr_tipolocal TEXT,
            descr_subtipolocal TEXT,
            cidade TEXT,
            bairro TEXT,
            cep TEXT,
            logradouro TEXT,
            numero_logradouro TEXT,
            latitude REAL,
            longitude REAL,
            marca_objeto TEXT,
            flag_bloqueio TEXT,
            flag_desbloqueio TEXT,
            mes_registro_bo INTEGER,
            ano_registro_bo INTEGER
        )
    """)
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_latlon ON crimes(latitude, longitude)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_cidade ON crimes(cidade)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_rubrica ON crimes(rubrica)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_data ON crimes(data_ocorrencia_bo)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_bairro ON crimes(bairro)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_crimes_marca ON crimes(marca_objeto)")
    conn.commit()
    return conn

def clean_val(v):
    """Convert None and 'NULL' strings to None for SQLite NULL."""
    if v is None or str(v) == 'NULL':
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if hasattr(v, 'strftime') and not isinstance(v, datetime):
        return v.strftime("%H:%M:%S")
    return v

def import_data():
    print(f"Opening {XLSX_PATH}...")
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
    ws = wb['CELULAR_2025']
    
    conn = create_db()
    c = conn.cursor()
    
    # Check if data already imported
    c.execute("SELECT COUNT(*) FROM crimes")
    count = c.fetchone()[0]
    if count > 0:
        print(f"Database already has {count} rows. Skipping import.")
        conn.close()
        wb.close()
        return
    
    # Build column list string
    cols_str = ",".join([c.lower() for c in WANTED_COLS])
    placeholders = ",".join(["?"] * len(WANTED_COLS))
    insert_sql = f"INSERT INTO crimes ({cols_str}) VALUES ({placeholders})"
    
    col_indices = {}
    batch = []
    batch_size = 5000
    total = 0
    skipped = 0
    
    LAT_IDX_IN_WANTED = WANTED_COLS.index('LATITUDE')  # 30
    LON_IDX_IN_WANTED = WANTED_COLS.index('LONGITUDE')  # 31
    
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = list(row)
            for col_name in WANTED_COLS:
                col_indices[col_name] = header.index(col_name)
            continue
        
        # Extract wanted columns
        vals = []
        for col_name in WANTED_COLS:
            vals.append(clean_val(row[col_indices[col_name]]))
        
        # Skip rows without valid coordinates
        lat = vals[LAT_IDX_IN_WANTED]
        lon = vals[LON_IDX_IN_WANTED]
        if lat is None or lon is None or float(lat) == 0.0 or float(lon) == 0.0:
            skipped += 1
            continue
        
        batch.append(vals)
        total += 1
        
        if len(batch) >= batch_size:
            c.executemany(insert_sql, batch)
            conn.commit()
            batch = []
            print(f"  Imported {total} rows...")
    
    if batch:
        c.executemany(insert_sql, batch)
        conn.commit()
    
    print(f"\nDone! Imported {total} rows with valid coordinates. Skipped {skipped} rows without coords.")
    conn.close()
    wb.close()

if __name__ == "__main__":
    import_data()
